"""Сборка отчётов prNN.md -> docx/ПР-NN.docx (титульный лист, ГОСТ-подобное оформление).

Требуется pandoc и python-docx. Запуск из корня репозитория: python docs/reports/build.py
"""
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "docx"
REFERENCE = HERE / "reference.docx"  # генерируется, в git не хранится

PROJECT = "Проектирование и разработка информационной системы «Видеостудия/продакшн — управление проектами»"
STUDENT = "Никита ____________"
GROUP = "________"
TEACHER = "Макиевский С. Е."
YEAR = "2026"
PAGE_BREAK = '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```'


def set_font(style, name="Times New Roman", size=14, bold=None):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def build_reference():
    OUT.mkdir(exist_ok=True)
    subprocess.run(["pandoc", "-o", str(REFERENCE), "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(REFERENCE)
    for section in doc.sections:
        section.page_height, section.page_width = Cm(29.7), Cm(21)
        section.left_margin, section.right_margin = Cm(3), Cm(1.5)
        section.top_margin = section.bottom_margin = Cm(2)
    by_name = {s.name: s for s in doc.styles}

    class _Styles:
        def __getitem__(self, name):
            return by_name[name]

    styles = _Styles()
    for name in ("Normal", "Body Text", "First Paragraph"):
        st = styles[name]
        set_font(st)
        pf = st.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(1.25)
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before, pf.space_after = Pt(0), Pt(6)
    compact = styles["Compact"]
    set_font(compact, size=12)
    compact.paragraph_format.first_line_indent = Cm(0)
    compact.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    compact.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 14)):
        st = styles[name]
        set_font(st, size=size, bold=True)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(12), Pt(6)
        st.paragraph_format.first_line_indent = Cm(0)
    for name in ("Image Caption", "Table Caption", "Captioned Figure"):
        if name in by_name:
            st = styles[name]
            set_font(st, size=12, bold=False)
            st.font.italic = False
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            st.paragraph_format.first_line_indent = Cm(0)
    doc.save(REFERENCE)


def title_page(number, topic):
    def center(text, style="Title Center"):
        return f'::: {{custom-style="{style}"}}\n{text}\n:::\n'

    return "\n".join([
        center("Институт перспективных технологий и индустриального программирования"),
        center("Кафедра индустриального программирования"),
        center("&nbsp;"), center("&nbsp;"), center("&nbsp;"),
        center("**ОТЧЁТ**", "Title Big"),
        center(f"**по практической работе №{number}**"),
        center(f"«{topic}»"),
        center("по дисциплине «Создание программного обеспечения»"),
        center("&nbsp;"),
        center(f"Тема проекта: {PROJECT}"),
        center("&nbsp;"), center("&nbsp;"), center("&nbsp;"),
        center(f"Выполнил: студент группы {GROUP}  \n{STUDENT}", "Title Right"),
        center("&nbsp;"),
        center(f"Проверил: преподаватель  \n{TEACHER}", "Title Right"),
        center("&nbsp;"),
        center(YEAR),
        PAGE_BREAK,
    ])


def prepare_markdown(path):
    text = path.read_text(encoding="utf-8")
    meta, body = re.match(r"---\n(.*?)\n---\n(.*)", text, re.S).groups()
    number = int(re.search(r"number:\s*(\d+)", meta).group(1))
    topic = re.search(r'topic:\s*"(.*)"', meta).group(1)

    counter = 0

    def figure(m):
        nonlocal counter
        counter += 1
        return f"![Рисунок {counter} — {m.group(1)}]({m.group(2)}){{width=16cm}}"

    body = re.sub(r"!\[(.*?)\]\((.*?)\)", figure, body)
    return number, topic, title_page(number, topic) + "\n" + body


def add_table_borders(doc):
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "000000")
            borders.append(el)
        tbl_pr.append(borders)
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True


def fit_tables(doc, total_cm=16.5):
    """Растягивает таблицы на ширину текста, ширина столбцов пропорциональна объёму текста."""
    for table in doc.tables:
        cols = len(table.columns)
        weight = [0] * cols
        for row in table.rows:
            for i, cell in enumerate(row.cells[:cols]):
                weight[i] = max(weight[i], min(len(cell.text), 60))
        longest_word = [0] * cols
        for row in table.rows:
            for i, cell in enumerate(row.cells[:cols]):
                words = cell.text.split() or [""]
                longest_word[i] = max(longest_word[i], min(max(len(w) for w in words), 22))
        min_cm = [0.24 * n + 0.45 for n in longest_word]
        weight = [max(w, 4) for w in weight]
        cm = [total_cm * w / sum(weight) for w in weight]
        for _ in range(cols):  # поднимаем узкие столбцы до минимума за счёт широких
            short = [i for i in range(cols) if cm[i] < min_cm[i]]
            if not short:
                break
            for i in short:
                cm[i] = min_cm[i]
            fixed = sum(cm[i] for i in range(cols) if cm[i] <= min_cm[i])
            rest = [i for i in range(cols) if cm[i] > min_cm[i]]
            rest_weight = sum(weight[i] for i in rest) or 1
            for i in rest:
                cm[i] = (total_cm - fixed) * weight[i] / rest_weight
        widths = [Cm(c) for c in cm]
        tbl_pr = table._tbl.tblPr
        for tag in ("w:tblW", "w:tblLayout"):
            old = tbl_pr.find(qn(tag))
            if old is not None:
                tbl_pr.remove(old)
        tbl_w = OxmlElement("w:tblW")
        tbl_w.set(qn("w:w"), str(int(total_cm * 567)))
        tbl_w.set(qn("w:type"), "dxa")
        tbl_pr.append(tbl_w)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tbl_pr.append(layout)
        grid = table._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[i].cm * 567)))
        for row in table.rows:
            for i, cell in enumerate(row.cells[:cols]):
                cell.width = widths[i]


def style_title_paragraphs(doc):
    for p in doc.paragraphs:
        name = p.style.name
        if name.startswith("Title "):
            pf = p.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT if name == "Title Right" else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(20 if name == "Title Big" else 14)


def add_page_numbers(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.text = text
        run._r.append(el)
    doc.sections[0].different_first_page_header_footer = True


def build(path):
    number, topic, md = prepare_markdown(path)
    target = OUT / f"ПР-{number:02d} {topic.replace(':', '.').replace('/', '-')}.docx"
    subprocess.run(
        ["pandoc", "-f", "markdown", "-t", "docx", "--reference-doc", str(REFERENCE),
         "--resource-path", str(HERE), "-o", str(target)],
        input=md.encode("utf-8"), check=True,
    )
    doc = Document(target)
    for st in doc.styles:
        if st.name == "Source Code":
            set_font(st, name="Courier New", size=10)
            st.paragraph_format.first_line_indent = Cm(0)
            st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            st.font.size = Pt(9)
    add_table_borders(doc)
    fit_tables(doc)
    style_title_paragraphs(doc)
    add_page_numbers(doc)
    doc.save(target)
    return target


if __name__ == "__main__":
    build_reference()
    files = sorted(HERE.glob("pr*.md")) if len(sys.argv) == 1 else [Path(a) for a in sys.argv[1:]]
    for f in files:
        print(build(f))
